"""Two-sheet consolidated workbook (Test Cases, Issues) built from the main migration report."""
import os
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = '/Users/sharathsp/Documents/v3-services'
SOURCE = f'{ROOT}/reports/Dine-in-v3-Java-Migration-Report.xlsx'
OUT = f'{ROOT}/reports/Dine-in-v3-Test-Cases-and-Issues.xlsx'
SHOTS = f'{ROOT}/reports/screenshots'
THUMBS = '/tmp/consolidated_thumbs'
THUMB_SIZE = (150, 270)

HEADER_FILL = PatternFill('solid', fgColor='1F3864')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILLS = {
    'PASS': 'C6EFCE', 'FAIL': 'FFC7CE', 'NOT TESTED': 'EDEDED',
    'HIGH': 'FFC7CE', 'MEDIUM': 'FFEB9C', 'LOW': 'DDEBF7', 'INFO': 'EDEDED',
}

src = load_workbook(SOURCE)
rows = lambda name: [r for r in src[name].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]

# ---------------------------------------------------------------- screenshot mapping (java, ruby)
TC_SHOTS = {
    'E2E-02': ('01-welcome', '01-welcome'), 'E2E-03': ('02-simple-item-added', '02-simple-item-added'),
    'E2E-04': ('03-customised-item-added', '03-customised-item-added'), 'E2E-05': ('04-cart', '04-cart'),
    'E2E-06': ('05-order-in-kitchen', '05-order-in-kitchen'), 'E2E-07': ('06-guest-joined', None),
    'E2E-08': ('07-bill', '07-bill'), 'E2E-09': ('08-card-form', '08-card-form'), 'E2E-10': ('09-paid', '09-paid'),
    'E2E-11': ('09-paid', '09-paid'), 'E2E-12': ('10-paid-after-reload', '10-paid-after-reload'),
    'E2E-13': ('11-feedback-submitted', '11-feedback-submitted'), 'E2E-18': ('12-start-new-order', '11-feedback-submitted'),
    'ITEM-01': ('04-cart', '04-cart'), 'ITEM-02': ('05-order-in-kitchen', '05-order-in-kitchen'), 'ITEM-03': ('07-bill', '07-bill'),
    'LR-01': ('live-1-add-onion-samosa', 'live-1-add-onion-samosa'), 'LR-02': ('diagnostic-3-linked-item-selected', 'live-3-linked-item-selected'),
    'LR-03': ('live-2-often-bought-with', 'live-2-often-bought-with'), 'LR-04': ('live-4-cart-bestseller', 'live-4-cart-bestseller'),
    'LR-05': ('diagnostic-1-add-onion-samosa', None), 'BASE-01': (None, '09-paid'), 'BASE-02': (None, '07-bill-split'), 'BASE-03': (None, '09-paid'),
}
ISSUE_SHOTS = {
    'ISS-01': ('05-order-in-kitchen', '05-order-in-kitchen'), 'ISS-02': ('07-bill', '07-bill'), 'ISS-03': ('09-paid', '09-paid'),
    'ISS-04': ('live-1-add-onion-samosa', 'live-1-add-onion-samosa'), 'ISS-10': ('07-bill', '07-bill-split'),
}
STAGE_SHOTS = {'1.': ('04-cart', '04-cart'), '2.': ('05-order-in-kitchen', '05-order-in-kitchen'), '3.': ('07-bill', '07-bill'),
               '4.': ('09-paid', '09-paid'), '5.': ('09-paid', '09-paid')}
AREA = {'API contract (Java)': 'API', 'Ruby vs Java parity': 'API', 'End-to-end (Java)': 'UI flow', 'Item comparison': 'Order data',
        'Linked & recommended': 'UI flow', 'Ruby baseline': 'Control run'}
SCREEN = {
    'E2E-01': 'Precondition', 'E2E-02': 'Welcome sheet / Menu', 'E2E-03': 'Menu', 'E2E-04': 'Item customisation sheet', 'E2E-05': 'Cart',
    'E2E-06': 'Place your order', 'E2E-07': 'Join table (PIN)', 'E2E-08': 'Pay bill', 'E2E-09': 'Payment method / Telr',
    'E2E-10': 'Payment success', 'E2E-11': 'Payment success', 'E2E-12': 'Payment success (reopened)', 'E2E-13': 'Rating & feedback',
    'E2E-14': 'Table released', 'E2E-15': 'All screens', 'E2E-16': 'All screens', 'E2E-17': 'All screens', 'E2E-18': 'Start a new order',
    'LR-01': 'Complete your meal sheet', 'LR-02': 'Complete your meal sheet', 'LR-03': 'Often bought with', 'LR-04': 'Cart Bestseller',
    'LR-05': 'Complete your meal / Often bought with', 'LR-06': 'Menu API',
}

os.makedirs(THUMBS, exist_ok=True)


def thumb(folder, name):
    if not name:
        return None
    path = f'{SHOTS}/{folder}/{name}.png'
    if not os.path.exists(path):
        return None
    out = f'{THUMBS}/{folder}-{name}.png'
    if not os.path.exists(out):
        im = PILImage.open(path)
        im.thumbnail(THUMB_SIZE)
        im.save(out)
    return out


def write_sheet(ws, title, summary, headers, data, widths, status_cols, shot_cols):
    ws.append([title])
    ws.cell(1, 1).font = Font(bold=True, size=14, color='1F3864')
    for line in summary:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(headers)
    for c in ws[header_row]:
        c.fill, c.font, c.border = HEADER_FILL, Font(bold=True, color='FFFFFF'), BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for values, shots in data:
        ws.append(values + [''] * len(shot_cols))
        r = ws.max_row
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical='top')
            c.border = BORDER
            if c.column in status_cols and isinstance(c.value, str) and c.value.upper() in FILLS:
                c.fill = PatternFill('solid', fgColor=FILLS[c.value.upper()])
                c.font = Font(bold=True)
        height = 0
        for col, (folder, name) in zip(shot_cols, shots):
            img = thumb(folder, name)
            if img:
                xl = XLImage(img)
                ws.add_image(xl, f'{get_column_letter(col)}{r}')
                height = max(height, PILImage.open(img).height * 0.75 + 8)
            else:
                ws.cell(r, col).value = '—'
        if height:
            ws.row_dimensions[r].height = height
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(header_row + 1, 3)
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{ws.max_row}'


wb = Workbook()

# ---------------------------------------------------------------- Test Cases
tc_rows = rows('Functional Test Cases')
counts = Counter(r[6] for r in tc_rows)
tc_data = []
for n, (tid, area, case, steps, expected, actual, result, ref) in enumerate(tc_rows, 1):
    j, rb = TC_SHOTS.get(tid, (None, None))
    tc_data.append(([n, tid, AREA.get(area, area), area, SCREEN.get(tid, 'API' if AREA.get(area) == 'API' else ''), case, steps, expected, actual, result, ref or ''],
                    [('java', j), ('ruby', rb)]))
write_sheet(
    wb.active, 'Dine-in QR (qr_menu_v3) – Java migration – all test cases',
    [f'Total: {len(tc_rows)}   ·   Passed: {counts["PASS"]}   ·   Failed: {counts["FAIL"]}   ·   Not tested: {counts.get("NOT TESTED", 0)}',
     'All API calls from qr_menu_v3 go to dine-in-service-staging.paytmcheckout.com (0 calls to the Ruby host). Failed cases are linked to the Issues sheet.'],
    ['S.No', 'Test ID', 'Type', 'Area', 'Screen / feature', 'Test case', 'Steps / input', 'Expected', 'Actual', 'Result', 'Issue ref',
     'Screenshot – Java (qr_menu_v3)', 'Screenshot – Ruby (qr_menu)'],
    tc_data, [6, 9, 11, 18, 20, 38, 32, 38, 42, 10, 16, 22, 22], status_cols=(10,), shot_cols=(12, 13))
wb.active.title = 'Test Cases'

# ---------------------------------------------------------------- Issues (defects + findings + data/response mismatches)
issue_data = []
failed_by_issue = {}
for r in tc_rows:
    if r[6] == 'FAIL' and r[7]:
        for ref in str(r[7]).replace('ISS-01..03', 'ISS-01, ISS-02, ISS-03').replace(',', ' ').split():
            if ref.startswith('ISS-'):
                failed_by_issue.setdefault(ref, []).append(r[0])

TYPE = {'CONFIRMED REGRESSION': 'Defect – confirmed regression', 'NEEDS DEV CHECK': 'Finding – needs dev check', 'ENVIRONMENT ONLY': 'Environment note'}
for iid, sev, cls, title, api, steps, expected, actual, impact, evidence in rows('Issues'):
    issue_data.append(([iid, TYPE.get(cls, cls), sev, title, api, steps, expected, actual, impact, ', '.join(failed_by_issue.get(iid, [])) or evidence],
                       [('java', ISSUE_SHOTS.get(iid, (None, None))[0]), ('ruby', ISSUE_SHOTS.get(iid, (None, None))[1])]))

n = 0
for section, item, detail, jv, rv, match, note in rows('Item Details (Java vs Ruby)'):
    if match != 'DIFFERENT':
        continue
    n += 1
    ref = 'ISS-11' if 'line count' in detail else 'ISS-03' if 'Paid successfully' in detail or 'Success screen' in detail else 'ISS-01 / ISS-02'
    sev = 'LOW' if ref == 'ISS-11' else 'MEDIUM' if ref == 'ISS-03' else 'HIGH'
    shots = STAGE_SHOTS.get(section[:2], (None, None))
    issue_data.append(([f'DM-{n:02d}', 'Data mismatch – order / bill / payment', sev, f'{item}: {detail}', section.split('. ', 1)[-1],
                        'Same items in both flows: Dahi ke Kebab 23.80; Paneer Dum Biryani 27.00 + Biryani Salan Dip 1.50 + Special Spiced Dates 3.00',
                        f'Ruby: {rv}', f'Java: {jv}', note or f'Related to {ref}', ref], [('java', shots[0]), ('ruby', shots[1])]))

n = 0
for endpoint, field, count, rv, jv, ref in rows('Ruby vs Java Diffs'):
    if not count:
        continue
    n += 1
    shots = ('live-2-often-bought-with', 'live-2-often-bought-with') if 'recommended' in field else \
            ('live-1-add-onion-samosa', 'live-1-add-onion-samosa') if 'linked' in field else (None, None)
    sev = 'HIGH' if ('recommended' in field or 'linked' in field) else 'LOW'
    issue_data.append(([f'RM-{n:02d}', 'API response mismatch (Ruby vs Java)', sev, f'{field} ({count} occurrence{"s" if count > 1 else ""})', endpoint,
                        'Send the same request to Ruby and Java and compare responses field by field', f'Ruby: {rv}', f'Java: {jv}',
                        'Clients relying on this field behave differently', ref], [('java', shots[0]), ('ruby', shots[1])]))

n = 0
for feature, where, source, rv, jv, diag, result, *_ in rows('Linked & Recommended'):
    if result != 'FAIL':
        continue
    n += 1
    shot = {'"Complete your meal" bottom sheet on Add': 'live-1-add-onion-samosa', 'Linked items listed in the sheet': 'live-1-add-onion-samosa',
            'Add a linked item from the sheet': 'live-3-linked-item-selected', '"Often bought with" section after adding': 'live-2-often-bought-with',
            'Recommended items listed': 'live-2-often-bought-with'}.get(feature)
    issue_data.append(([f'LR-M{n:02d}', 'UI / data mismatch – linked & recommended', 'HIGH', f'{feature} – {where}', source,
                        'Menu → Soups & Snacks → Add on Onion Samosa (then Aloo & Cheese Samosa + 7 UP)', f'Ruby: {rv}', f'Java: {jv}',
                        f'With Ruby ids injected into the Java response the v3 UI shows: {diag} → Java API defect', 'ISS-04'],
                       [('java', shot if shot != 'live-3-linked-item-selected' else 'diagnostic-3-linked-item-selected'), ('ruby', shot)]))

sev_counts = Counter(r[0][2] for r in issue_data)
type_counts = Counter(r[0][1] for r in issue_data)
write_sheet(
    wb.create_sheet('Issues'), 'Dine-in QR (qr_menu_v3) – issues, findings and data / response mismatches',
    [f'Total rows: {len(issue_data)}   ·   High: {sev_counts["HIGH"]}   ·   Medium: {sev_counts["MEDIUM"]}   ·   Low: {sev_counts["LOW"]}   ·   Info: {sev_counts["INFO"]}',
     '   ·   '.join(f'{t}: {c}' for t, c in type_counts.items()),
     'Expected = Ruby (old API) or correct behaviour · Actual = Java (new API). Top-level issues ISS-xx; DM/RM/LR-M rows are the individual mismatches behind them.'],
    ['ID', 'Type', 'Severity', 'Title / field', 'API / screen', 'Steps to reproduce', 'Expected (Ruby / correct)', 'Actual (Java)', 'Impact / note',
     'Failed test cases / ref', 'Screenshot – Java (qr_menu_v3)', 'Screenshot – Ruby (qr_menu)'],
    issue_data, [9, 24, 10, 38, 28, 38, 38, 42, 36, 18, 22, 22], status_cols=(3,), shot_cols=(11, 12))

wb.save(OUT)
print(OUT, '| test cases:', len(tc_rows), dict(counts), '| issue rows:', len(issue_data), dict(sev_counts))
